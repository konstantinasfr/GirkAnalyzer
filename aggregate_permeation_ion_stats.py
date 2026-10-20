#!/usr/bin/env python3
"""
aggregate_permeation_ion_stats.py
=================================
Cross-run summary of the permeation_ion_* pipeline, at BOTH levels:

  PER-RUN   each run's ions are averaged into one row, then the runs are
            averaged. Every run counts equally regardless of how many ions it
            had. Good for "do the runs agree", and the historical behaviour of
            this script.

  PER-ION   every permeating ion is one data point, pooled across all runs.
            No averaging before the plot. Good for "what does an ion do", and
            for comparing channels on the ion distribution itself.

Both are produced in one call. -noion skips the per-ion section, -nobox skips
the per-run boxplots, so the old output set is still reachable.

INPUT LAYOUT
------------
    <base_dir>/                     e.g. .../girk_analyser_results/G12
        RUN1/permeation_ion_table/permeation_ion_table_<unit>.csv   (per-ion)
        RUN1/permeation_ion_table/stats_<unit>/residue_percentages_by_group.csv
        RUN1/permeation_ion_table/stats_<unit>/mobility_<unit>.csv
        RUN2/...
        permeation_summary/permeation_analysis.xlsx   (optional)

Several channel folders can be given at once. Each gets its own per-run
outputs; the per-ion figures pool them so G12 and G2 sit side by side:

    python3 aggregate_permeation_ion_stats.py .../G12 .../G2 -stats

PER-ION COLUMNS USED
--------------------
    cavity_<unit>    the ion's cavity residence  -> y of the residence boxplot
    res_<LABEL>      its contact TIME with that residue, same unit
    frames_window    the measured window, same unit
    start_censored / incomplete_range    truncation flags

The res_ columns are times, not percentages, so the contact boxplot divides
each by the ion's own residence. That is what makes a 5 ns ion and a 170 ns
ion comparable; a raw contact time would just re-plot residence.

CENSORED IONS are dropped by default: their residence was cut off by the start
or end of the trajectory, so the recorded value is a lower bound, not a
measurement. The number dropped is printed per run. -keep_censored overrides.

The `run` column inside the per-ion table is unreliable (some runs write 0),
so run names come from the folder.

OUTPUTS (in <base_dir>/permeation_ion_aggregate/)
-------------------------------------------------
  per-run
    residue_percentages_matrix.csv   one row per run, 24 residue x group columns
    residue_percentages_mean.csv     the column means across runs
    mobility_matrix.csv              one row per run
    mobility_mean.csv                the column means across runs
    residue_pct_all_runs.png         grouped barplot of the pooled means
    residue_pct_spread.png           per-run scatter over the pooled means
    residue_pct_box.png              per-run distribution as boxplots
    mobility_all_runs.png            pooled mobility, stacked
    permeation_counts_box.png        permeations per run
    permeation_counts_per_run.csv

  per-ion (pooled across the channels given, written to the FIRST channel's
  folder, or to -out)
    pooled_ions.csv                  every ion used, with channel and run
    residence_box.png                y = that ion's residence; dot = one ion
    residue_contact_box.png          y = % of that ion's cavity time in contact
    residence_summary.csv
    residue_contact_summary.csv
    channel_tests.csv                with -stats

TESTING THE PER-ION BOXES
-------------------------
With -stats the channel comparison is printed twice, and the two answer
different questions:

  ion-level      Mann-Whitney over all ions. This is what the dots show, but
                 ions from one trajectory are not independent — same cavity,
                 same starting structure, same neighbours. It will look more
                 significant than the data warrant.
  run-clustered  Mann-Whitney over per-run medians, n = number of runs. This
                 treats the trajectory as the unit of replication and is the
                 number to quote in a paper.

If the two disagree sharply, trust the clustered one. Per-residue tests are
run-clustered with Benjamini-Hochberg FDR across residues.

PLOT SIZING
-----------
Every text element has an explicit font size in the FONT SIZES block below, so
nothing inherits matplotlib's default 10 pt and disappears when the figure is
shrunk into a slide. Scale them all with -fontscale (figure sizes scale too).

USAGE
-----
    python3 aggregate_permeation_ion_stats.py .../G12
    python3 aggregate_permeation_ion_stats.py .../G12 .../G2 -stats
    python3 aggregate_permeation_ion_stats.py <base_dir> -stat median -unit ns
    python3 aggregate_permeation_ion_stats.py <base_dir> -weight ions
    python3 aggregate_permeation_ion_stats.py <base_dir> -fontscale 1.2
    python3 aggregate_permeation_ion_stats.py <base_dir> -by run      # per-ion boxes per run
    python3 aggregate_permeation_ion_stats.py <base_dir> -noion       # old outputs only
"""

import argparse
import csv
import re
import sys
from pathlib import Path

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

SCRIPT_VERSION = "2026-10-13b  (per-run + per-ion in one script)"

MAIN_GROUPS = ["below_Q1", "Q1_Q3", "above_Q3"]
GROUP_LABEL = {"below_Q1": "Below Q1 (fast)",
               "Q1_Q3": "Q1-Q3 (typical)",
               "above_Q3": "Above Q3 (slow)"}
GROUP_COLOR = {"below_Q1": "#27AE60",
               "Q1_Q3": "#2980B9",
               "above_Q3": "#E67E22"}

CHANNEL_COLOR = {"G12": "#2980B9", "G2": "#E67E22",
                 "G12_ML": "#27AE60", "G12_GAT": "#8E44AD"}

RUN_PATTERN = re.compile(r"^RUN(\d+)$", re.IGNORECASE)
RES_PREFIX = "res_"

# ─────────────────────────────────────────────────────────────────────────────
#  FONT SIZES — all multiplied by -fontscale at run time.
# ─────────────────────────────────────────────────────────────────────────────
FS_TITLE   = 23   # plot titles
FS_AXLABEL = 21   # axis labels
FS_TICK    = 19   # tick labels
FS_VALUE   = 16   # numbers drawn on the bars
FS_LEGEND  = 18

DPI = 200
BAR_EDGE_LW = 1.0

# Base figure sizes, in inches, scaled by -fontscale
FIG_RESIDUE_BARS = (16.0, 7.0)
FIG_SPREAD       = (16.0, 7.2)
FIG_RESIDUE_BOX  = (17.0, 7.4)
FIG_MOBILITY     = (12.0, 7.4)
FIG_COUNTS_BOX   = (9.5, 7.0)
FIG_RESIDENCE    = (10.0, 7.0)
FIG_ION_RESIDUE  = (17.0, 7.4)

# Count columns taken from permeation_analysis.xlsx, with display labels
COUNT_COLS = [("hbc_permeations", "HBC\npermeations"),
              ("full_permeations_unique", "Full permeations\n(unique)"),
              ("full_permeations_total", "Full permeations\n(total)")]


# ─────────────────────────────────────────────────────────────────────────────
#  Small helpers
# ─────────────────────────────────────────────────────────────────────────────

def _num(v, default=None):
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return default


def _truthy(v):
    return str(v).strip().lower() in ("true", "1", "yes")


def _finish(ax, fs):
    ax.tick_params(axis="both", labelsize=FS_TICK * fs)
    ax.grid(True, alpha=0.25, axis="y")
    ax.spines[["top", "right"]].set_visible(False)


def write_rows(rows, path, fieldnames=None):
    if not rows:
        print(f"  [warn] nothing to write to {Path(path).name}")
        return
    if fieldnames is None:
        fieldnames = []
        for r in rows:
            for k in r:
                if k not in fieldnames:
                    fieldnames.append(k)
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    print(f"  Saved: {Path(path).name}  ({len(rows)} rows)")


def weighted_mean(pairs):
    """pairs = [(value, weight)]; ignores missing values."""
    vals = [(v, w) for v, w in pairs if v is not None and w]
    if not vals:
        return ""
    total_w = sum(w for _, w in vals)
    if total_w == 0:
        return ""
    return round(sum(v * w for v, w in vals) / total_w, 3)


def summarise(label, vals):
    v = np.asarray(vals, float)
    return dict(group=label, n_ions=int(v.size),
                median=round(float(np.median(v)), 3),
                mean=round(float(v.mean()), 3),
                q1=round(float(np.percentile(v, 25)), 3),
                q3=round(float(np.percentile(v, 75)), 3),
                minimum=round(float(v.min()), 3),
                maximum=round(float(v.max()), 3))


# ─────────────────────────────────────────────────────────────────────────────
#  Discovery
# ─────────────────────────────────────────────────────────────────────────────

def find_runs(base_dir, unit):
    """Every RUN* folder, with whichever of the three input files it has."""
    found = []
    for sub in base_dir.iterdir():
        if not sub.is_dir():
            continue
        m = RUN_PATTERN.match(sub.name)
        if not m:
            continue
        table_dir = sub / "permeation_ion_table"
        stats_dir = table_dir / f"stats_{unit}"
        pct_file = stats_dir / "residue_percentages_by_group.csv"
        mob_file = stats_dir / f"mobility_{unit}.csv"
        ion_file = table_dir / f"permeation_ion_table_{unit}.csv"
        found.append(dict(run=sub.name, index=int(m.group(1)),
                          stats_dir=stats_dir,
                          pct=pct_file if pct_file.is_file() else None,
                          mob=mob_file if mob_file.is_file() else None,
                          ions=ion_file if ion_file.is_file() else None))
    found.sort(key=lambda d: d["index"])
    return found


# ─────────────────────────────────────────────────────────────────────────────
#  Permeation counts
# ─────────────────────────────────────────────────────────────────────────────

def load_permeation_counts(base_dir):
    """
    Read Per_Run_Details from permeation_summary/permeation_analysis.xlsx.
    Returns {RUN_NAME: {col: value}} or {} if unavailable.
    """
    xlsx = base_dir / "permeation_summary" / "permeation_analysis.xlsx"
    if not xlsx.is_file():
        print(f"  [info] {xlsx} not found — permeation counts will be blank")
        return {}
    try:
        import openpyxl
    except ImportError:
        print("  [warn] openpyxl not installed — permeation counts will be blank "
              "(pip install openpyxl)")
        return {}

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if "Per_Run_Details" not in wb.sheetnames:
        print(f"  [warn] no Per_Run_Details sheet in {xlsx.name}")
        return {}
    ws = wb["Per_Run_Details"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    out = {}
    for r in rows[1:]:
        if not r or r[0] is None:
            continue
        rec = dict(zip(header, r))
        out[str(rec["Run"]).upper()] = rec
    print(f"  Loaded permeation counts for {len(out)} runs from {xlsx.name}")
    return out


def count_columns(counts, run_name):
    """The permeation columns appended to the end of every per-run row."""
    rec = counts.get(run_name.upper(), {})
    return {
        "hbc_permeations": rec.get("HBC Permeations", ""),
        "full_permeations_unique": rec.get("Full Permeations Unique", ""),
        "full_permeations_total": rec.get("Full Permeations Total", ""),
        "total_frames": rec.get("Total Frames", ""),
    }


# ─────────────────────────────────────────────────────────────────────────────
#  Reading one run — per-run level
# ─────────────────────────────────────────────────────────────────────────────

def read_pct_file(path, stat):
    """
    residue_percentages_by_group.csv -> {(residue, group): value}, plus
    {group: n_ions}. Only the three main groups are kept; the extreme sets are
    dropped here because they overlap the quartile bands.
    """
    key = "mean_pct" if stat == "mean" else "median_pct"
    values, n_ions = {}, {}
    with open(path) as f:
        for row in csv.DictReader(f):
            g = row["group"]
            if g not in MAIN_GROUPS:
                continue
            values[(row["residue"], g)] = float(row[key])
            n_ions[g] = int(row["n_ions"])
    return values, n_ions


def read_mob_file(path, stat):
    """mobility_<unit>.csv -> {(kind, group): pct} for kind in any/no."""
    any_key = "mean_pct_any" if stat == "mean" else "median_pct_any"
    no_key = "mean_pct_no" if stat == "mean" else "median_pct_no"
    out = {}
    with open(path) as f:
        for row in csv.DictReader(f):
            g = row["group"]
            if g not in MAIN_GROUPS:
                continue
            out[("any", g)] = float(row[any_key])
            out[("no", g)] = float(row[no_key])
    return out


def residue_order(runs, stat):
    """Residue labels in the order the first available run lists them."""
    for r in runs:
        if r["pct"]:
            with open(r["pct"]) as f:
                seen = []
                for row in csv.DictReader(f):
                    if row["group"] in MAIN_GROUPS and row["residue"] not in seen:
                        seen.append(row["residue"])
            if seen:
                return seen
    return []


# ─────────────────────────────────────────────────────────────────────────────
#  Reading one run — per-ion level
# ─────────────────────────────────────────────────────────────────────────────

def load_ions(path, run_name, channel, unit, keep_censored):
    """One dict per ion. Returns (ions, n_dropped)."""
    ycol = f"cavity_{unit}"
    ions, dropped = [], 0
    with open(path, newline="") as f:
        rows = list(csv.DictReader(f))
    if not rows:
        return [], 0
    if ycol not in rows[0]:
        print(f"    [warn] {run_name}: no '{ycol}' column in the per-ion table "
              f"— skipped")
        return [], 0

    res_cols = [c for c in rows[0] if c.startswith(RES_PREFIX)]
    for r in rows:
        y = _num(r.get(ycol))
        if y is None or y <= 0:
            continue
        censored = (_truthy(r.get("start_censored"))
                    or _truthy(r.get("incomplete_range")))
        if censored and not keep_censored:
            dropped += 1
            continue
        ion = dict(
            channel_type=(r.get("channel_type") or channel).strip() or channel,
            run=run_name,
            ion_rank=r.get("ion_rank", ""),
            ion_id=r.get("ion_id", ""),
            residence=y,
            window=_num(r.get("frames_window"), y),
            any_contact=_num(r.get("frames_any_contact")),
            no_contact=_num(r.get("frames_no_contact")),
            censored=censored,
        )
        for c in res_cols:
            ion[c] = _num(r.get(c), 0.0)
        ions.append(ion)
    return ions, dropped


def ion_residue_labels(ions):
    """res_ columns in file order, label only."""
    if not ions:
        return []
    return [c[len(RES_PREFIX):] for c in ions[0] if c.startswith(RES_PREFIX)]


def contact_pct(ion, label, denom):
    """That ion's contact time with one residue, as % of its own residence."""
    d = ion["window"] if denom == "window" else ion["residence"]
    if not d:
        return None
    return 100.0 * ion.get(RES_PREFIX + label, 0.0) / d


# ─────────────────────────────────────────────────────────────────────────────
#  Statistics (per-ion section)
# ─────────────────────────────────────────────────────────────────────────────

def mannwhitney(a, b):
    a, b = np.asarray(a, float), np.asarray(b, float)
    if a.size < 3 or b.size < 3:
        return float("nan"), float("nan"), f"too few points (n={a.size}, {b.size})"
    try:
        from scipy import stats
        u, p = stats.mannwhitneyu(a, b, alternative="two-sided")
        return float(u), float(p), ""
    except ImportError:
        return float("nan"), float("nan"), "scipy not installed"
    except Exception as e:
        return float("nan"), float("nan"), str(e)


def bh_fdr(pvals):
    p = np.asarray(pvals, dtype=float)
    ok = ~np.isnan(p)
    q = np.full(p.shape, np.nan)
    if not ok.any():
        return q
    sub = p[ok]
    m = sub.size
    order = np.argsort(sub)
    adj = sub[order] * m / np.arange(1, m + 1)
    adj = np.clip(np.minimum.accumulate(adj[::-1])[::-1], 0, 1)
    out = np.empty(m)
    out[order] = adj
    q[ok] = out
    return q


def stars(q):
    if q != q:
        return ""
    for thr, s in [(0.001, "***"), (0.01, "**"), (0.05, "*")]:
        if q < thr:
            return s
    return "ns"


def _fmt_p(v):
    if v != v:
        return ""
    return f"{v:.2e}" if v < 1e-4 else round(float(v), 6)


def per_run_medians(ions, value_fn):
    """{run: median over that run's ions} — the clustered view."""
    by_run = {}
    for ion in ions:
        v = value_fn(ion)
        if v is not None:
            by_run.setdefault(ion["run"], []).append(v)
    return {k: float(np.median(v)) for k, v in by_run.items() if v}


# ─────────────────────────────────────────────────────────────────────────────
#  Plots — per-run level
# ─────────────────────────────────────────────────────────────────────────────

def plot_pooled_residue_bars(residues, means, n_runs, out_path, channel, stat,
                             fs=1.0):
    x = np.arange(len(residues))
    width = 0.8 / len(MAIN_GROUPS)
    fig, ax = plt.subplots(figsize=(FIG_RESIDUE_BARS[0] * fs,
                                    FIG_RESIDUE_BARS[1] * fs))
    for i, g in enumerate(MAIN_GROUPS):
        vals = [means.get((r, g), np.nan) for r in residues]
        ax.bar(x + i * width - 0.4 + width / 2, vals, width=width * 0.92,
               color=GROUP_COLOR[g], alpha=0.85, edgecolor="black",
               linewidth=BAR_EDGE_LW, label=GROUP_LABEL[g])
    ax.set_xticks(x)
    ax.set_xticklabels(residues, rotation=45, ha="right", fontsize=FS_TICK * fs)
    ax.set_ylabel(f"{stat.capitalize()} % of cavity time in contact",
                  fontsize=FS_AXLABEL * fs)
    ax.set_title(f"{channel} — residue contact profile by cavity-residence group\n"
                 f"(pooled over {n_runs} runs)",
                 fontsize=FS_TITLE * fs, fontweight="bold")
    # legend outside the axes so it never sits on top of a tall bar
    ax.legend(fontsize=FS_LEGEND * fs, loc="upper left",
              bbox_to_anchor=(1.01, 1.0), borderaxespad=0)
    _finish(ax, fs)
    plt.tight_layout()
    plt.savefig(out_path, dpi=DPI, bbox_inches="tight")
    plt.close()
    print(f"  Saved: {Path(out_path).name}")


def plot_spread(residues, per_run, means, out_path, channel, fs=1.0):
    """Pooled means as bars, with each run's own value as a dot on top."""
    x = np.arange(len(residues))
    width = 0.8 / len(MAIN_GROUPS)
    rng = np.random.default_rng(0)
    fig, ax = plt.subplots(figsize=(FIG_SPREAD[0] * fs, FIG_SPREAD[1] * fs))
    for i, g in enumerate(MAIN_GROUPS):
        pos = x + i * width - 0.4 + width / 2
        vals = [means.get((r, g), np.nan) for r in residues]
        ax.bar(pos, vals, width=width * 0.92, color=GROUP_COLOR[g], alpha=0.55,
               edgecolor="black", linewidth=BAR_EDGE_LW, label=GROUP_LABEL[g])
        for j, res in enumerate(residues):
            pts = [row["values"].get((res, g)) for row in per_run
                   if (res, g) in row["values"]]
            if pts:
                ax.scatter(pos[j] + rng.uniform(-width * 0.25, width * 0.25, len(pts)),
                           pts, s=26 * fs, color="black", alpha=0.45,
                           edgecolors="none", zorder=3)
    ax.set_xticks(x)
    ax.set_xticklabels(residues, rotation=45, ha="right", fontsize=FS_TICK * fs)
    ax.set_ylabel("% of cavity time in contact", fontsize=FS_AXLABEL * fs)
    ax.set_title(f"{channel} — per-run spread around the pooled mean\n"
                 f"(each dot is one run)",
                 fontsize=FS_TITLE * fs, fontweight="bold")
    ax.legend(fontsize=FS_LEGEND * fs, loc="upper left",
              bbox_to_anchor=(1.01, 1.0), borderaxespad=0)
    _finish(ax, fs)
    plt.tight_layout()
    plt.savefig(out_path, dpi=DPI, bbox_inches="tight")
    plt.close()
    print(f"  Saved: {Path(out_path).name}")


def plot_residue_box(residues, per_run, out_path, channel, fs=1.0):
    """
    One box per residue x residence group, over the RUNS. The bar plots show
    only where the centre sits; these boxes show how far the runs disagree.
    """
    if len(per_run) < 3:
        print(f"  [skip] residue boxplot needs >=3 runs, have {len(per_run)}")
        return
    x = np.arange(len(residues))
    width = 0.8 / len(MAIN_GROUPS)
    rng = np.random.default_rng(0)
    fig, ax = plt.subplots(figsize=(FIG_RESIDUE_BOX[0] * fs,
                                    FIG_RESIDUE_BOX[1] * fs))

    for i, g in enumerate(MAIN_GROUPS):
        data, pos = [], []
        for j, res in enumerate(residues):
            pts = [row["values"][(res, g)] for row in per_run
                   if (res, g) in row["values"]]
            if not pts:
                continue
            data.append(pts)
            pos.append(j + i * width - 0.4 + width / 2)
        if not data:
            continue
        bp = ax.boxplot(data, positions=pos, widths=width * 0.85,
                        patch_artist=True, showfliers=False,
                        medianprops=dict(color="green", linewidth=2.0 * fs),
                        showmeans=True, meanline=True,
                        meanprops=dict(color="red", linestyle="--",
                                       linewidth=1.6 * fs),
                        whiskerprops=dict(linewidth=1.2 * fs),
                        capprops=dict(linewidth=1.2 * fs))
        for box in bp["boxes"]:
            box.set_facecolor(GROUP_COLOR[g])
            box.set_alpha(0.6)
            box.set_edgecolor("black")
            box.set_linewidth(BAR_EDGE_LW)
        for p, pts in zip(pos, data):
            ax.scatter(p + rng.uniform(-width * 0.18, width * 0.18, len(pts)),
                       pts, s=22 * fs, color="black", alpha=0.40,
                       edgecolors="none", zorder=4)
        ax.plot([], [], color=GROUP_COLOR[g], linewidth=10 * fs, alpha=0.6,
                label=GROUP_LABEL[g])

    ax.plot([], [], color="green", linewidth=2.0 * fs, label="Median")
    ax.plot([], [], color="red", linestyle="--", linewidth=1.6 * fs, label="Mean")

    ax.set_xticks(x)
    ax.set_xticklabels(residues, rotation=45, ha="right", fontsize=FS_TICK * fs)
    ax.set_xlim(-0.6, len(residues) - 0.4)
    ax.set_ylabel("% of cavity time in contact", fontsize=FS_AXLABEL * fs)
    ax.set_title(f"{channel} — run-to-run distribution per residue\n"
                 f"(one box per residence group, {len(per_run)} runs, "
                 f"each dot one run)",
                 fontsize=FS_TITLE * fs, fontweight="bold")
    ax.legend(fontsize=FS_LEGEND * fs, loc="upper left",
              bbox_to_anchor=(1.01, 1.0), borderaxespad=0)
    _finish(ax, fs)
    plt.tight_layout()
    plt.savefig(out_path, dpi=DPI, bbox_inches="tight")
    plt.close()
    print(f"  Saved: {Path(out_path).name}")


def plot_permeation_counts_box(count_rows, out_path, channel, fs=1.0):
    """
    Permeations per run, as boxplots — the flux distribution of this channel.
    Each dot is a run, so a high-flux outlier trajectory stays visible instead
    of being averaged away. Returns the per-type values for the CSV.
    """
    series, labels = [], []
    for key, label in COUNT_COLS:
        vals = [r[key] for r in count_rows
                if isinstance(r[key], (int, float)) and r[key] != ""]
        if vals:
            series.append([float(v) for v in vals])
            labels.append(label)
    if not series:
        print("  [skip] permeation counts boxplot — no counts available "
              "(permeation_analysis.xlsx missing or unreadable)")
        return None
    if len(series[0]) < 3:
        print(f"  [skip] permeation counts boxplot needs >=3 runs, "
              f"have {len(series[0])}")
        return None

    rng = np.random.default_rng(0)
    fig, ax = plt.subplots(figsize=(max(FIG_COUNTS_BOX[0], 3.4 * len(series)) * fs,
                                    FIG_COUNTS_BOX[1] * fs))
    pos = np.arange(1, len(series) + 1)
    bp = ax.boxplot(series, positions=pos, widths=0.55,
                    patch_artist=True, showfliers=False,
                    medianprops=dict(color="green", linewidth=2.2 * fs),
                    showmeans=True, meanline=True,
                    meanprops=dict(color="red", linestyle="--",
                                   linewidth=1.8 * fs),
                    whiskerprops=dict(linewidth=1.3 * fs),
                    capprops=dict(linewidth=1.3 * fs))
    for box in bp["boxes"]:
        box.set_facecolor("#2980B9")
        box.set_alpha(0.6)
        box.set_edgecolor("black")
        box.set_linewidth(BAR_EDGE_LW)
    for p, vals in zip(pos, series):
        ax.scatter(p + rng.uniform(-0.14, 0.14, len(vals)), vals,
                   s=46 * fs, color="black", alpha=0.45, edgecolors="none",
                   zorder=4)

    ax.plot([], [], color="green", linewidth=2.2 * fs, label="Median")
    ax.plot([], [], color="red", linestyle="--", linewidth=1.8 * fs, label="Mean")

    ax.set_xticks(pos)
    ax.set_xticklabels(labels, fontsize=FS_TICK * fs)
    ax.set_ylabel("Permeations per run", fontsize=FS_AXLABEL * fs)
    ax.set_ylim(bottom=0)
    ax.set_title(f"{channel} — permeation events per run\n"
                 f"({len(series[0])} runs, each dot one run)",
                 fontsize=FS_TITLE * fs, fontweight="bold")
    ax.legend(fontsize=FS_LEGEND * fs, loc="upper left",
              bbox_to_anchor=(1.01, 1.0), borderaxespad=0)
    _finish(ax, fs)
    plt.tight_layout()
    plt.savefig(out_path, dpi=DPI, bbox_inches="tight")
    plt.close()
    print(f"  Saved: {Path(out_path).name}")
    return dict(zip([k for k, _ in COUNT_COLS], series))


def plot_pooled_mobility(mob_means, n_runs, out_path, channel, fs=1.0):
    groups = [g for g in MAIN_GROUPS if ("any", g) in mob_means]
    if not groups:
        return
    x = np.arange(len(groups))
    any_p = [mob_means[("any", g)] for g in groups]
    no_p = [mob_means[("no", g)] for g in groups]

    fig, ax = plt.subplots(figsize=(FIG_MOBILITY[0] * fs, FIG_MOBILITY[1] * fs))
    ax.bar(x, any_p, width=0.6, color="#2980B9", alpha=0.85,
           edgecolor="black", linewidth=BAR_EDGE_LW,
           label="In contact with a residue")
    ax.bar(x, no_p, width=0.6, bottom=any_p, color="#BDC3C7", alpha=0.9,
           edgecolor="black", linewidth=BAR_EDGE_LW, label="Touching nothing")
    for i, (a, n) in enumerate(zip(any_p, no_p)):
        ax.text(i, a / 2, f"{a:.0f}%", ha="center", va="center",
                fontsize=FS_VALUE * fs, color="white", fontweight="bold")
        ax.text(i, a + n / 2, f"{n:.0f}%", ha="center", va="center",
                fontsize=FS_VALUE * fs, color="black")
    ax.set_xticks(x)
    # two-line labels: the one-line version collides at these font sizes
    ax.set_xticklabels([GROUP_LABEL[g].replace(" (", "\n(") for g in groups],
                       fontsize=FS_TICK * fs)
    ax.set_ylabel("Mean % of cavity time", fontsize=FS_AXLABEL * fs)
    # the two segments sum to 100 by construction, so that is the ceiling
    ax.set_ylim(0, 100)
    ax.set_yticks([0, 20, 40, 60, 80, 100])
    ax.set_title(f"{channel} — ion mobility in the cavity\n"
                 f"(pooled over {n_runs} runs)",
                 fontsize=FS_TITLE * fs, fontweight="bold", pad=60 * fs)
    # legend above the bars, between them and the title, rather than at the
    # side: the bars reach 100% so there is no room for it inside the axes
    ax.legend(fontsize=FS_LEGEND * fs, loc="lower center",
              bbox_to_anchor=(0.5, 1.005), ncol=2, framealpha=0.9,
              borderaxespad=0)
    _finish(ax, fs)
    plt.tight_layout()
    plt.savefig(out_path, dpi=DPI, bbox_inches="tight")
    plt.close()
    print(f"  Saved: {Path(out_path).name}")


# ─────────────────────────────────────────────────────────────────────────────
#  Plots — per-ion level
# ─────────────────────────────────────────────────────────────────────────────

def plot_ion_residence(groups, out_path, unit, by, logy, fs=1.0, color_of=None):
    """One box per group, one dot per ION, y = that ion's residence."""
    labels = [k for k, v in groups.items() if v]
    if not labels:
        return []
    data = [groups[k] for k in labels]
    rng = np.random.default_rng(0)

    width = max(FIG_RESIDENCE[0], 2.6 * len(labels)) * fs
    fig, ax = plt.subplots(figsize=(width, FIG_RESIDENCE[1] * fs))
    pos = np.arange(1, len(labels) + 1)
    bp = ax.boxplot(data, positions=pos, widths=0.55, patch_artist=True,
                    showfliers=False,
                    medianprops=dict(color="green", linewidth=2.2 * fs),
                    showmeans=True, meanline=True,
                    meanprops=dict(color="red", linestyle="--",
                                   linewidth=1.8 * fs))
    for box, lab in zip(bp["boxes"], labels):
        box.set_facecolor((color_of or {}).get(lab, "#2980B9"))
        box.set_alpha(0.6)
        box.set_edgecolor("black")
        box.set_linewidth(BAR_EDGE_LW)
    for w in bp["whiskers"] + bp["caps"]:
        w.set_linewidth(1.3 * fs)

    for p, vals in zip(pos, data):
        v = np.asarray(vals, float)
        ax.scatter(p + rng.uniform(-0.16, 0.16, v.size), v,
                   s=40 * fs, color="black", alpha=0.40, edgecolors="none",
                   zorder=4)

    ax.plot([], [], color="green", linewidth=2.2 * fs, label="Median")
    ax.plot([], [], color="red", linestyle="--", linewidth=1.8 * fs, label="Mean")

    ax.set_xticks(pos)
    ax.set_xticklabels([f"{l}\n(n = {len(groups[l])} ions)" for l in labels],
                       fontsize=FS_TICK * fs)
    ax.set_ylabel(f"Cavity residence ({'ns' if unit == 'ns' else 'frames'})",
                  fontsize=FS_AXLABEL * fs)
    if logy:
        ax.set_yscale("log")
    ax.set_title(f"Cavity residence per permeating ion — by {by}\n"
                 f"(each dot is one ion)" + ("   [log scale]" if logy else ""),
                 fontsize=FS_TITLE * fs, fontweight="bold")
    ax.legend(fontsize=FS_LEGEND * fs, loc="upper left",
              bbox_to_anchor=(1.01, 1.0), borderaxespad=0)
    _finish(ax, fs)
    plt.tight_layout()
    plt.savefig(out_path, dpi=DPI, bbox_inches="tight")
    plt.close()
    print(f"  Saved: {Path(out_path).name}")
    return labels


def assign_groups(ions, mode):
    """
    Put each ion in a residence band: below Q1 (fast), Q1-Q3 (typical),
    above Q3 (slow).

    mode "pooled"  quartiles from all this channel's ions together. One set of
                   thresholds, so "fast" means the same thing for every ion in
                   the figure.
    mode "perrun"  quartiles from the ion's own run, which is what
                   permeation_ion_stats.py does upstream. Matches the bar plot,
                   but "fast" is then relative to each trajectory: a fast ion in
                   a slow run can be slower than a slow ion in a fast run.

    Returns {ion_index_in_list: group}, plus the thresholds used (for printing).
    """
    out, info = {}, {}
    if mode == "perrun":
        by_run = {}
        for k, ion in enumerate(ions):
            by_run.setdefault(ion["run"], []).append(k)
        for run, idxs in by_run.items():
            vals = np.array([ions[k]["residence"] for k in idxs], float)
            if vals.size < 4:
                # too few ions to split into quartiles sensibly
                for k in idxs:
                    out[k] = "Q1_Q3"
                continue
            q1, q3 = np.percentile(vals, 25), np.percentile(vals, 75)
            info[run] = (q1, q3)
            for k in idxs:
                v = ions[k]["residence"]
                out[k] = ("below_Q1" if v < q1 else
                          "above_Q3" if v > q3 else "Q1_Q3")
    else:
        vals = np.array([i["residence"] for i in ions], float)
        q1, q3 = np.percentile(vals, 25), np.percentile(vals, 75)
        info["pooled"] = (q1, q3)
        for k, ion in enumerate(ions):
            v = ion["residence"]
            out[k] = ("below_Q1" if v < q1 else
                      "above_Q3" if v > q3 else "Q1_Q3")
    return out, info


def plot_ion_residue_by_group(labels, ions, groups_of, out_path, channel,
                              denom, mode, fs=1.0):
    """
    The same layout as residue_pct_all_runs.png — residue on x, one bar per
    residence group — but every dot is an ION, and the bar is the mean over
    those ions rather than a mean of per-run means. A green line marks the
    median, which on these strongly right-skewed contact distributions usually
    sits well below the bar top; the gap between the two is itself the signal
    that a few long-contact ions are carrying the mean.
    """
    x = np.arange(len(labels))
    width = 0.8 / len(MAIN_GROUPS)
    rng = np.random.default_rng(0)
    fig, ax = plt.subplots(figsize=(max(FIG_ION_RESIDUE[0], 2.0 * len(labels)) * fs,
                                    FIG_ION_RESIDUE[1] * fs))

    n_in_group = {g: sum(1 for v in groups_of.values() if v == g)
                  for g in MAIN_GROUPS}
    for i, g in enumerate(MAIN_GROUPS):
        members = [ions[k] for k, gg in groups_of.items() if gg == g]
        if not members:
            continue
        heights, medians, pos = [], [], []
        points = []
        for j, lab in enumerate(labels):
            vals = [contact_pct(m, lab, denom) for m in members]
            vals = [v for v in vals if v is not None]
            if not vals:
                continue
            v = np.asarray(vals, float)
            heights.append(float(v.mean()))
            medians.append(float(np.median(v)))
            pos.append(j + i * width - 0.4 + width / 2)
            points.append(v)
        if not heights:
            continue
        # bar = mean over the ions of this group, matching the bar plot this
        # figure is modelled on; the dots are the ions it was computed from
        ax.bar(pos, heights, width=width * 0.92, color=GROUP_COLOR[g],
               alpha=0.80, edgecolor="black", linewidth=BAR_EDGE_LW, zorder=2)
        for p, m in zip(pos, medians):
            ax.plot([p - width * 0.42, p + width * 0.42], [m, m],
                    color="green", linewidth=2.0 * fs, zorder=5)
        for p, v in zip(pos, points):
            ax.scatter(p + rng.uniform(-width * 0.20, width * 0.20, v.size), v,
                       s=18 * fs, color="black", alpha=0.32, edgecolors="none",
                       zorder=4)
        ax.plot([], [], color=GROUP_COLOR[g], linewidth=10 * fs, alpha=0.8,
                label=f"{GROUP_LABEL[g]} — {n_in_group[g]} ions")

    ax.plot([], [], color="green", linewidth=2.0 * fs, label="Median")
    ax.scatter([], [], s=18 * fs, color="black", alpha=0.5, label="One ion")

    ax.set_xticks(x)
    ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=FS_TICK * fs)
    ax.set_xlim(-0.6, len(labels) - 0.4)
    ax.set_ylabel("% of the ion's cavity time in contact",
                  fontsize=FS_AXLABEL * fs)
    ax.set_ylim(bottom=0)
    qtxt = ("quartiles over all ions of this channel" if mode == "pooled"
            else "quartiles within each run")
    ax.set_title(f"{channel} — residue contact by cavity-residence group\n"
                 f"(bars = mean over ions, each dot one ion, {len(ions)} ions; "
                 f"{qtxt})",
                 fontsize=FS_TITLE * fs, fontweight="bold")
    ax.legend(fontsize=FS_LEGEND * fs, loc="upper left",
              bbox_to_anchor=(1.01, 1.0), borderaxespad=0)
    _finish(ax, fs)
    plt.tight_layout()
    plt.savefig(out_path, dpi=DPI, bbox_inches="tight")
    plt.close()
    print(f"  Saved: {Path(out_path).name}")


def plot_ion_residue_by_group_box(labels, ions, groups_of, out_path, channel,
                                  denom, mode, fs=1.0):
    """
    Same figure as plot_ion_residue_by_group, drawn as boxes instead of bars.

    The two are worth keeping side by side: the bars match the historical
    figure and are easy to read, the boxes show that most of these contact
    distributions are strongly right-skewed — a median near zero under a mean
    of several percent means a few ions are carrying the bar.
    """
    x = np.arange(len(labels))
    width = 0.8 / len(MAIN_GROUPS)
    rng = np.random.default_rng(0)
    fig, ax = plt.subplots(figsize=(max(FIG_ION_RESIDUE[0], 2.0 * len(labels)) * fs,
                                    FIG_ION_RESIDUE[1] * fs))

    n_in_group = {g: sum(1 for v in groups_of.values() if v == g)
                  for g in MAIN_GROUPS}
    for i, g in enumerate(MAIN_GROUPS):
        members = [ions[k] for k, gg in groups_of.items() if gg == g]
        if not members:
            continue
        data, pos = [], []
        for j, lab in enumerate(labels):
            vals = [contact_pct(m, lab, denom) for m in members]
            vals = [v for v in vals if v is not None]
            if not vals:
                continue
            data.append(vals)
            pos.append(j + i * width - 0.4 + width / 2)
        if not data:
            continue
        bp = ax.boxplot(data, positions=pos, widths=width * 0.85,
                        patch_artist=True, showfliers=False,
                        medianprops=dict(color="green", linewidth=2.0 * fs),
                        showmeans=True, meanline=True,
                        meanprops=dict(color="red", linestyle="--",
                                       linewidth=1.5 * fs),
                        whiskerprops=dict(linewidth=1.2 * fs),
                        capprops=dict(linewidth=1.2 * fs))
        for box in bp["boxes"]:
            box.set_facecolor(GROUP_COLOR[g])
            box.set_alpha(0.6)
            box.set_edgecolor("black")
            box.set_linewidth(BAR_EDGE_LW)
        for p, vals in zip(pos, data):
            v = np.asarray(vals, float)
            ax.scatter(p + rng.uniform(-width * 0.18, width * 0.18, v.size), v,
                       s=18 * fs, color="black", alpha=0.32, edgecolors="none",
                       zorder=4)
        ax.plot([], [], color=GROUP_COLOR[g], linewidth=10 * fs, alpha=0.6,
                label=f"{GROUP_LABEL[g]} — {n_in_group[g]} ions")

    ax.plot([], [], color="green", linewidth=2.0 * fs, label="Median")
    ax.plot([], [], color="red", linestyle="--", linewidth=1.5 * fs, label="Mean")

    ax.set_xticks(x)
    ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=FS_TICK * fs)
    ax.set_xlim(-0.6, len(labels) - 0.4)
    ax.set_ylabel("% of the ion's cavity time in contact",
                  fontsize=FS_AXLABEL * fs)
    ax.set_ylim(bottom=0)
    qtxt = ("quartiles over all ions of this channel" if mode == "pooled"
            else "quartiles within each run")
    ax.set_title(f"{channel} — residue contact by cavity-residence group\n"
                 f"(each dot is one ion, {len(ions)} ions; {qtxt})",
                 fontsize=FS_TITLE * fs, fontweight="bold")
    ax.legend(fontsize=FS_LEGEND * fs, loc="upper left",
              bbox_to_anchor=(1.01, 1.0), borderaxespad=0)
    _finish(ax, fs)
    plt.tight_layout()
    plt.savefig(out_path, dpi=DPI, bbox_inches="tight")
    plt.close()
    print(f"  Saved: {Path(out_path).name}")


def plot_ion_mobility_by_group(ions, groups_of, out_path, channel, mode, fs=1.0):
    """
    The per-ion form of mobility_all_runs.png.

    The stacked bar cannot carry dots — a stack is two numbers that sum to 100,
    not a distribution — so the same information is drawn as the distribution of
    the one free number: the % of each ion's cavity time spent in contact with
    ANY of the residues. "Touching nothing" is its complement, so the right-hand
    axis labels it directly and no information is lost.
    """
    data, pos, counts = [], [], []
    for i, g in enumerate(MAIN_GROUPS):
        vals = []
        for k, gg in groups_of.items():
            if gg != g:
                continue
            ion = ions[k]
            d = ion["window"] or ion["residence"]
            if ion["any_contact"] is None or not d:
                continue
            vals.append(100.0 * ion["any_contact"] / d)
        if not vals:
            continue
        data.append(vals)
        pos.append(len(data))
        counts.append((g, len(vals)))
    if not data:
        print(f"  [skip] {channel}: no frames_any_contact column in the "
              f"per-ion table — mobility boxplot skipped")
        return

    rng = np.random.default_rng(0)
    fig, ax = plt.subplots(figsize=(FIG_MOBILITY[0] * fs, FIG_MOBILITY[1] * fs))
    bp = ax.boxplot(data, positions=pos, widths=0.55, patch_artist=True,
                    showfliers=False,
                    medianprops=dict(color="green", linewidth=2.2 * fs),
                    showmeans=True, meanline=True,
                    meanprops=dict(color="red", linestyle="--",
                                   linewidth=1.8 * fs),
                    whiskerprops=dict(linewidth=1.3 * fs),
                    capprops=dict(linewidth=1.3 * fs))
    for box, (g, _n) in zip(bp["boxes"], counts):
        box.set_facecolor(GROUP_COLOR[g])
        box.set_alpha(0.6)
        box.set_edgecolor("black")
        box.set_linewidth(BAR_EDGE_LW)
    for p, vals in zip(pos, data):
        v = np.asarray(vals, float)
        ax.scatter(p + rng.uniform(-0.16, 0.16, v.size), v,
                   s=40 * fs, color="black", alpha=0.40, edgecolors="none",
                   zorder=4)

    ax.plot([], [], color="green", linewidth=2.2 * fs, label="Median")
    ax.plot([], [], color="red", linestyle="--", linewidth=1.8 * fs, label="Mean")

    ax.set_xticks(pos)
    ax.set_xticklabels([f"{GROUP_LABEL[g]}".replace(" (", "\n(")
                        + f"\n{n} ions" for g, n in counts],
                       fontsize=FS_TICK * fs)
    ax.set_ylabel("% of the ion's cavity time in contact",
                  fontsize=FS_AXLABEL * fs)
    ax.set_ylim(0, 100)
    ax.set_yticks([0, 20, 40, 60, 80, 100])

    # the complement, so the "touching nothing" reading is still available
    ax2 = ax.twinx()
    ax2.set_ylim(100, 0)
    ax2.set_yticks([0, 20, 40, 60, 80, 100])
    ax2.set_ylabel("% touching nothing", fontsize=FS_AXLABEL * fs)
    ax2.tick_params(axis="y", labelsize=FS_TICK * fs)
    ax2.spines[["top"]].set_visible(False)

    qtxt = ("quartiles over all ions of this channel" if mode == "pooled"
            else "quartiles within each run")
    ax.set_title(f"{channel} — ion mobility in the cavity\n"
                 f"(each dot is one ion, {len(ions)} ions; {qtxt})",
                 fontsize=FS_TITLE * fs, fontweight="bold", pad=55 * fs)
    ax.legend(fontsize=FS_LEGEND * fs, loc="lower center",
              bbox_to_anchor=(0.5, 1.005), ncol=2, framealpha=0.9,
              borderaxespad=0)
    _finish(ax, fs)
    plt.tight_layout()
    plt.savefig(out_path, dpi=DPI, bbox_inches="tight")
    plt.close()
    print(f"  Saved: {Path(out_path).name}")


def plot_ion_residue_contact(labels, by_channel, out_path, fs=1.0):
    """One box per residue per channel, one dot per ION."""
    channels = [c for c in by_channel if by_channel[c]]
    if not channels or not labels:
        return
    x = np.arange(len(labels))
    width = 0.8 / len(channels)
    rng = np.random.default_rng(0)

    fig, ax = plt.subplots(figsize=(max(FIG_ION_RESIDUE[0], 2.0 * len(labels)) * fs,
                                    FIG_ION_RESIDUE[1] * fs))
    n_ions = {}
    for i, ch in enumerate(channels):
        data, pos = [], []
        for j, lab in enumerate(labels):
            vals = by_channel[ch].get(lab, [])
            if not vals:
                continue
            data.append(vals)
            pos.append(j + i * width - 0.4 + width / 2)
            n_ions[ch] = max(n_ions.get(ch, 0), len(vals))
        if not data:
            continue
        bp = ax.boxplot(data, positions=pos, widths=width * 0.85,
                        patch_artist=True, showfliers=False,
                        medianprops=dict(color="green", linewidth=2.0 * fs),
                        showmeans=True, meanline=True,
                        meanprops=dict(color="red", linestyle="--",
                                       linewidth=1.5 * fs))
        for box in bp["boxes"]:
            box.set_facecolor(CHANNEL_COLOR.get(ch, "#777777"))
            box.set_alpha(0.6)
            box.set_edgecolor("black")
            box.set_linewidth(BAR_EDGE_LW)
        for w in bp["whiskers"] + bp["caps"]:
            w.set_linewidth(1.2 * fs)
        for p, vals in zip(pos, data):
            v = np.asarray(vals, float)
            ax.scatter(p + rng.uniform(-width * 0.18, width * 0.18, v.size), v,
                       s=20 * fs, color="black", alpha=0.35, edgecolors="none",
                       zorder=4)
        ax.plot([], [], color=CHANNEL_COLOR.get(ch, "#777777"),
                linewidth=10 * fs, alpha=0.6,
                label=f"{ch} (n = {n_ions.get(ch, 0)} ions)")

    ax.plot([], [], color="green", linewidth=2.0 * fs, label="Median")
    ax.plot([], [], color="red", linestyle="--", linewidth=1.5 * fs, label="Mean")

    ax.set_xticks(x)
    ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=FS_TICK * fs)
    ax.set_xlim(-0.6, len(labels) - 0.4)
    ax.set_ylabel("% of the ion's cavity time in contact",
                  fontsize=FS_AXLABEL * fs)
    ax.set_ylim(bottom=0)
    ax.set_title("Residue contact per permeating ion\n"
                 "(each dot is one ion, pooled over all runs)",
                 fontsize=FS_TITLE * fs, fontweight="bold")
    ax.legend(fontsize=FS_LEGEND * fs, loc="upper left",
              bbox_to_anchor=(1.01, 1.0), borderaxespad=0)
    _finish(ax, fs)
    plt.tight_layout()
    plt.savefig(out_path, dpi=DPI, bbox_inches="tight")
    plt.close()
    print(f"  Saved: {Path(out_path).name}")


# ─────────────────────────────────────────────────────────────────────────────
#  Per-run section, one channel
# ─────────────────────────────────────────────────────────────────────────────

def run_per_run_section(base_dir, runs, channel, out_dir, args, fs):
    have_pct = [r for r in runs if r["pct"]]
    have_mob = [r for r in runs if r["mob"]]
    print(f"  {len(runs)} RUN folders; {len(have_pct)} with percentages, "
          f"{len(have_mob)} with mobility")
    for r in runs:
        if not r["pct"]:
            print(f"    [missing] {r['run']}: no residue_percentages_by_group.csv")
        elif not r["mob"]:
            print(f"    [missing] {r['run']}: no mobility_{args.unit}.csv")

    if not have_pct:
        print("  [skip] per-run section: no residue_percentages_by_group.csv "
              "in any run")
        return

    counts = load_permeation_counts(base_dir)
    residues = residue_order(have_pct, args.stat)
    if not residues:
        print("  [skip] per-run section: no residues found")
        return
    print(f"  {len(residues)} residues: {residues}")

    pct_cols = [f"{res}_{g}" for res in residues for g in MAIN_GROUPS]
    n_cols = [f"n_ions_{g}" for g in MAIN_GROUPS]
    cnt_cols = ["hbc_permeations", "full_permeations_unique",
                "full_permeations_total", "total_frames"]

    pct_rows, per_run = [], []
    for r in have_pct:
        values, n_ions = read_pct_file(r["pct"], args.stat)
        row = dict(run=r["run"], channel_type=channel, stat=args.stat)
        for res in residues:
            for g in MAIN_GROUPS:
                row[f"{res}_{g}"] = values.get((res, g), "")
        for g in MAIN_GROUPS:
            row[f"n_ions_{g}"] = n_ions.get(g, "")
        row.update(count_columns(counts, r["run"]))
        pct_rows.append(row)
        per_run.append(dict(run=r["run"], values=values, n_ions=n_ions))

    write_rows(pct_rows, out_dir / "residue_percentages_matrix.csv",
               ["run", "channel_type", "stat"] + pct_cols + n_cols + cnt_cols)

    means = {}
    mean_row = {"run": f"MEAN_OF_{len(pct_rows)}_RUNS", "channel_type": channel,
                "stat": args.stat}
    for res in residues:
        for g in MAIN_GROUPS:
            pairs = [(row["values"].get((res, g)),
                      row["n_ions"].get(g, 1) if args.weight == "ions" else 1)
                     for row in per_run]
            m = weighted_mean(pairs)
            means[(res, g)] = float(m) if m != "" else np.nan
            mean_row[f"{res}_{g}"] = m
    for g in MAIN_GROUPS:
        mean_row[f"n_ions_{g}"] = sum(row["n_ions"].get(g, 0) for row in per_run)
    for c in cnt_cols:
        vals = [r[c] for r in pct_rows if r[c] != ""]
        mean_row[c] = sum(vals) if vals else ""
    write_rows([mean_row], out_dir / "residue_percentages_mean.csv",
               ["run", "channel_type", "stat"] + pct_cols + n_cols + cnt_cols)

    print("\n  Pooled percentages:")
    print(f"    {'residue':<12}" + "".join(f"{GROUP_LABEL[g][:10]:>12}"
                                           for g in MAIN_GROUPS))
    for res in residues:
        line = f"    {res:<12}"
        for g in MAIN_GROUPS:
            v = means.get((res, g), np.nan)
            line += f"{v:>12.2f}" if v == v else f"{'-':>12}"
        print(line)

    mob_means = {}
    if have_mob:
        mob_cols = [f"{k}_contact_{g}" for g in MAIN_GROUPS for k in ("any", "no")]
        mob_rows, mob_per_run = [], []
        for r in have_mob:
            vals = read_mob_file(r["mob"], args.stat)
            row = dict(run=r["run"], channel_type=channel, stat=args.stat)
            for g in MAIN_GROUPS:
                row[f"any_contact_{g}"] = vals.get(("any", g), "")
                row[f"no_contact_{g}"] = vals.get(("no", g), "")
            row.update(count_columns(counts, r["run"]))
            mob_rows.append(row)
            mob_per_run.append(vals)
        write_rows(mob_rows, out_dir / "mobility_matrix.csv",
                   ["run", "channel_type", "stat"] + mob_cols + cnt_cols)

        mrow = {"run": f"MEAN_OF_{len(mob_rows)}_RUNS", "channel_type": channel,
                "stat": args.stat}
        for g in MAIN_GROUPS:
            for kind in ("any", "no"):
                vals = [v[(kind, g)] for v in mob_per_run if (kind, g) in v]
                if vals:
                    mob_means[(kind, g)] = float(np.mean(vals))
                    mrow[f"{kind}_contact_{g}"] = round(float(np.mean(vals)), 3)
                else:
                    mrow[f"{kind}_contact_{g}"] = ""
        for c in cnt_cols:
            vals = [r[c] for r in mob_rows if r[c] != ""]
            mrow[c] = sum(vals) if vals else ""
        write_rows([mrow], out_dir / "mobility_mean.csv",
                   ["run", "channel_type", "stat"] + mob_cols + cnt_cols)

        print("\n  Pooled mobility:")
        for g in MAIN_GROUPS:
            if ("any", g) in mob_means:
                print(f"    {g:<14} in contact {mob_means[('any', g)]:>5.1f}%   "
                      f"free {mob_means[('no', g)]:>5.1f}%")
    else:
        print("\n  [info] no mobility files found — re-run permeation_ion_table.py "
              "(2026-09-26c+) and permeation_ion_stats.py to generate them")

    print()
    plot_pooled_residue_bars(residues, means, len(pct_rows),
                             out_dir / "residue_pct_all_runs.png",
                             channel, args.stat, fs=fs)
    plot_spread(residues, per_run, means,
                out_dir / "residue_pct_spread.png", channel, fs=fs)
    if mob_means:
        plot_pooled_mobility(mob_means, len(have_mob),
                             out_dir / "mobility_all_runs.png", channel, fs=fs)

    if not args.nobox:
        plot_residue_box(residues, per_run,
                         out_dir / "residue_pct_box.png", channel, fs=fs)
        series = plot_permeation_counts_box(
            pct_rows, out_dir / "permeation_counts_box.png", channel, fs=fs)
        if series:
            rows = [dict(run=r["run"], channel_type=channel,
                         **{k: r[k] for k, _ in COUNT_COLS},
                         total_frames=r["total_frames"])
                    for r in pct_rows]
            write_rows(rows, out_dir / "permeation_counts_per_run.csv",
                       ["run", "channel_type"] + [k for k, _ in COUNT_COLS]
                       + ["total_frames"])
            print("\n  Permeations per run:")
            for key, label in COUNT_COLS:
                v = np.array(series.get(key, []), dtype=float)
                if v.size:
                    print(f"    {label.replace(chr(10), ' '):<28} "
                          f"median {np.median(v):>6.1f}   mean {v.mean():>6.1f}   "
                          f"range {v.min():.0f}-{v.max():.0f}   n={v.size} runs")


# ─────────────────────────────────────────────────────────────────────────────
#  Per-ion section, all channels together
# ─────────────────────────────────────────────────────────────────────────────

def run_per_ion_section(all_ions, total_dropped, out_dir, args, fs):
    channels = []
    for ion in all_ions:
        if ion["channel_type"] not in channels:
            channels.append(ion["channel_type"])

    print(f"\n  {len(all_ions)} ions, {len(channels)} channel(s)"
          + (f", {total_dropped} censored ions dropped" if total_dropped else ""))

    labels = ion_residue_labels(all_ions)
    print(f"  {len(labels)} residues: {labels}")

    rows = []
    for ion in all_ions:
        row = dict(channel_type=ion["channel_type"], run=ion["run"],
                   ion_rank=ion["ion_rank"], ion_id=ion["ion_id"],
                   residence=round(ion["residence"], 4),
                   censored=ion["censored"])
        for lab in labels:
            v = contact_pct(ion, lab, args.denom)
            row[f"pct_{lab}"] = round(v, 3) if v is not None else ""
        rows.append(row)
    write_rows(rows, out_dir / "pooled_ions.csv")

    # residence boxplot
    if args.by == "channel":
        groups = {ch: [i["residence"] for i in all_ions
                       if i["channel_type"] == ch] for ch in channels}
        color_of = {ch: CHANNEL_COLOR.get(ch, "#2980B9") for ch in channels}
    else:
        groups, color_of = {}, {}
        for ion in all_ions:
            key = (f"{ion['channel_type']} {ion['run']}" if len(channels) > 1
                   else ion["run"])
            groups.setdefault(key, []).append(ion["residence"])
            color_of[key] = CHANNEL_COLOR.get(ion["channel_type"], "#2980B9")

    used = plot_ion_residence(groups, out_dir / "residence_box.png", args.unit,
                              args.by, not args.linear, fs=fs, color_of=color_of)
    res_summary = [summarise(l, groups[l]) for l in used]
    write_rows(res_summary, out_dir / "residence_summary.csv")
    print("\n  Residence per ion:")
    for s in res_summary:
        print(f"    {s['group']:<18} n={s['n_ions']:<4} median={s['median']:<9} "
              f"mean={s['mean']:<9} range {s['minimum']}-{s['maximum']}")

    # residue contact boxplot
    by_channel = {}
    for ch in channels:
        d = {}
        for lab in labels:
            vals = [contact_pct(i, lab, args.denom) for i in all_ions
                    if i["channel_type"] == ch]
            vals = [v for v in vals if v is not None]
            if vals:
                d[lab] = vals
        by_channel[ch] = d
    plot_ion_residue_contact(labels, by_channel,
                             out_dir / "residue_contact_box.png", fs=fs)

    contact_summary = []
    for ch in channels:
        for lab in labels:
            if lab in by_channel[ch]:
                s = summarise(lab, by_channel[ch][lab])
                s["channel_type"] = ch
                contact_summary.append(s)
    write_rows(contact_summary, out_dir / "residue_contact_summary.csv",
               ["channel_type", "group", "n_ions", "median", "mean", "q1", "q3",
                "minimum", "maximum"])

    # the residue x residence-group layout of residue_pct_all_runs.png,
    # but one dot per ion — one figure per channel, since three groups x
    # two channels on one axis would be unreadable
    group_summary = []
    for ch in channels:
        ions_ch = [i for i in all_ions if i["channel_type"] == ch]
        if len(ions_ch) < 4:
            print(f"  [skip] {ch}: only {len(ions_ch)} ions, too few to split "
                  f"into residence quartiles")
            continue
        groups_of, info = assign_groups(ions_ch, args.quartiles)
        if args.quartiles == "pooled" and "pooled" in info:
            q1, q3 = info["pooled"]
            print(f"  {ch} residence quartiles: Q1={q1:.2f}, Q3={q3:.2f} "
                  f"{args.unit}")
        plot_ion_residue_by_group(
            labels, ions_ch, groups_of,
            out_dir / f"residue_contact_bars_by_group_{ch}.png",
            ch, args.denom, args.quartiles, fs=fs)
        plot_ion_residue_by_group_box(
            labels, ions_ch, groups_of,
            out_dir / f"residue_contact_box_by_group_{ch}.png",
            ch, args.denom, args.quartiles, fs=fs)
        plot_ion_mobility_by_group(
            ions_ch, groups_of, out_dir / f"mobility_box_by_group_{ch}.png",
            ch, args.quartiles, fs=fs)
        for g in MAIN_GROUPS:
            members = [ions_ch[k] for k, gg in groups_of.items() if gg == g]
            for lab in labels:
                vals = [contact_pct(m, lab, args.denom) for m in members]
                vals = [v for v in vals if v is not None]
                if vals:
                    s = summarise(lab, vals)
                    s["channel_type"] = ch
                    s["residence_group"] = g
                    group_summary.append(s)
    if group_summary:
        write_rows(group_summary,
                   out_dir / "residue_contact_by_group_summary.csv",
                   ["channel_type", "residence_group", "group", "n_ions",
                    "median", "mean", "q1", "q3", "minimum", "maximum"])

    # tests
    if args.stats and len(channels) >= 2:
        print("\n  === channel comparison (per-ion data) ===")
        print("  (ion-level ignores that ions from one run are not independent;"
              "\n   the run-clustered row is the one to quote)")
        test_rows = []
        for i in range(len(channels)):
            for j in range(i + 1, len(channels)):
                a_ch, b_ch = channels[i], channels[j]
                a_ions = [x for x in all_ions if x["channel_type"] == a_ch]
                b_ions = [x for x in all_ions if x["channel_type"] == b_ch]

                a_i = [x["residence"] for x in a_ions]
                b_i = [x["residence"] for x in b_ions]
                u, p, note = mannwhitney(a_i, b_i)
                test_rows.append(dict(
                    quantity="cavity residence", level="ion-level",
                    group_a=a_ch, group_b=b_ch, n_a=len(a_i), n_b=len(b_i),
                    median_a=round(float(np.median(a_i)), 3),
                    median_b=round(float(np.median(b_i)), 3),
                    statistic=round(u, 3) if u == u else "",
                    p_value=_fmt_p(p), q_value="", stars="", note=note))

                a_r = list(per_run_medians(a_ions, lambda x: x["residence"]).values())
                b_r = list(per_run_medians(b_ions, lambda x: x["residence"]).values())
                u2, p2, note2 = mannwhitney(a_r, b_r)
                test_rows.append(dict(
                    quantity="cavity residence", level="run-clustered",
                    group_a=a_ch, group_b=b_ch, n_a=len(a_r), n_b=len(b_r),
                    median_a=round(float(np.median(a_r)), 3) if a_r else "",
                    median_b=round(float(np.median(b_r)), 3) if b_r else "",
                    statistic=round(u2, 3) if u2 == u2 else "",
                    p_value=_fmt_p(p2), q_value="", stars="", note=note2))

                per_res = []
                for lab in labels:
                    a_rr = list(per_run_medians(
                        a_ions, lambda x, l=lab: contact_pct(x, l, args.denom)).values())
                    b_rr = list(per_run_medians(
                        b_ions, lambda x, l=lab: contact_pct(x, l, args.denom)).values())
                    uu, pp, nn = mannwhitney(a_rr, b_rr)
                    per_res.append(dict(
                        quantity=f"contact % {lab}", level="run-clustered",
                        group_a=a_ch, group_b=b_ch, n_a=len(a_rr), n_b=len(b_rr),
                        median_a=round(float(np.median(a_rr)), 3) if a_rr else "",
                        median_b=round(float(np.median(b_rr)), 3) if b_rr else "",
                        statistic=round(uu, 3) if uu == uu else "",
                        p_value=_fmt_p(pp), q_value="", stars="", note=nn))
                qs = bh_fdr([float(r["p_value"]) if r["p_value"] != "" else np.nan
                             for r in per_res])
                for r, q in zip(per_res, qs):
                    r["q_value"] = _fmt_p(q)
                    r["stars"] = stars(q)
                test_rows.extend(per_res)

        for t in test_rows:
            if t["quantity"] == "cavity residence":
                print(f"    {t['level']:<14} {t['group_a']} (med {t['median_a']}, "
                      f"n={t['n_a']}) vs {t['group_b']} (med {t['median_b']}, "
                      f"n={t['n_b']}): p={t['p_value']}"
                      + (f"  [{t['note']}]" if t["note"] else ""))
        print("\n    per-residue contact %, run-clustered, BH-FDR:")
        for t in test_rows:
            if t["quantity"].startswith("contact %"):
                print(f"      {t['quantity'][10:]:<12} "
                      f"{t['median_a']:>7} vs {t['median_b']:>7}  "
                      f"p={t['p_value']} q={t['q_value']} {t['stars']}")
        write_rows(test_rows, out_dir / "channel_tests.csv")
    elif args.stats:
        print("\n  [note] only one channel loaded — pass both channel folders "
              "in the same call for the comparison.")


# ─────────────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("base_dir", nargs="+",
                    help="Channel folder(s) holding RUN1, RUN2, ... "
                         "Pass G12 and G2 together for the per-ion comparison.")
    ap.add_argument("-unit", choices=["frames", "ns"], default="frames",
                    help="Which stats_<unit>/ and permeation_ion_table_<unit>.csv "
                         "to read (default frames)")
    ap.add_argument("-stat", choices=["mean", "median"], default="mean",
                    help="Which per-run statistic to pool (default mean)")
    ap.add_argument("-weight", choices=["run", "ions"], default="run",
                    help="PER-RUN section: run = every run counts equally "
                         "(default); ions = weight each run by its ion count")
    ap.add_argument("-channel", default=None,
                    help="Label for plot titles when ONE folder is given "
                         "(default: the folder name)")
    ap.add_argument("-out", default=None,
                    help="Write everything to this one directory instead of "
                         "each channel's own permeation_ion_aggregate/")
    ap.add_argument("-fontscale", "--fontscale", type=float, default=1.0,
                    help="Multiply every font size (and figure size) by this")
    ap.add_argument("-nobox", "--nobox", action="store_true",
                    help="Skip the per-run boxplots")
    ap.add_argument("-noion", "--noion", action="store_true",
                    help="Skip the per-ion section entirely")
    # per-ion options
    ap.add_argument("-by", choices=["channel", "run"], default="channel",
                    help="PER-ION residence boxes: one per channel (default) "
                         "or one per run")
    ap.add_argument("-denom", choices=["cavity", "window"], default="cavity",
                    help="PER-ION contact %%: divide by the ion's cavity "
                         "residence (default) or by frames_window")
    ap.add_argument("-quartiles", choices=["pooled", "perrun"], default="pooled",
                    help="PER-ION fast/typical/slow bands: quartiles over all "
                         "ions of a channel (default), or within each run as "
                         "permeation_ion_stats.py does upstream")
    ap.add_argument("-linear", action="store_true",
                    help="PER-ION residence plot on a linear y-axis "
                         "(default: log, since the distribution is skewed)")
    ap.add_argument("-keep_censored", action="store_true",
                    help="PER-ION: keep ions flagged start_censored / "
                         "incomplete_range (dropped by default)")
    ap.add_argument("-stats", action="store_true",
                    help="PER-ION: Mann-Whitney between channels, ion-level "
                         "and run-clustered")
    args = ap.parse_args()

    fs = args.fontscale
    dirs = [Path(d) for d in args.base_dir]
    for d in dirs:
        if not d.is_dir():
            sys.exit(f"ERROR: not a directory: {d}")

    print("=" * 70)
    print("AGGREGATE PERMEATION ION STATISTICS")
    print("=" * 70)
    print(f"  version: {SCRIPT_VERSION}")
    print(f"  unit={args.unit}  stat={args.stat}  weight={args.weight}  "
          f"by={args.by}  censored={'kept' if args.keep_censored else 'dropped'}")

    forced_out = Path(args.out).resolve() if args.out else None
    if forced_out:
        forced_out.mkdir(parents=True, exist_ok=True)

    all_ions, total_dropped, first_out = [], 0, None

    for d in dirs:
        channel = args.channel if (args.channel and len(dirs) == 1) else d.name
        out_dir = forced_out or (d / "permeation_ion_aggregate")
        out_dir.mkdir(parents=True, exist_ok=True)
        if first_out is None:
            first_out = out_dir

        print(f"\n{'=' * 70}\n  {channel}   ({d})\n{'=' * 70}")
        runs = find_runs(d, args.unit)
        if not runs:
            print(f"  [warn] no RUN* folders in {d}")
            continue

        run_per_run_section(d, runs, channel, out_dir, args, fs)

        if not args.noion:
            have_ions = [r for r in runs if r["ions"]]
            if not have_ions:
                print(f"\n  [info] no permeation_ion_table_{args.unit}.csv found "
                      f"— per-ion section skipped for {channel}")
            else:
                print(f"\n  Per-ion tables: {len(have_ions)} runs")
                for r in have_ions:
                    ions, dropped = load_ions(r["ions"], r["run"], channel,
                                              args.unit, args.keep_censored)
                    total_dropped += dropped
                    all_ions.extend(ions)
                    note = f", {dropped} censored dropped" if dropped else ""
                    print(f"    {r['run']}: {len(ions)} ions{note}")

    if not args.noion and all_ions:
        ion_out = forced_out or first_out
        print(f"\n{'=' * 70}\n  PER-ION SECTION  (every ion one data point)"
              f"\n{'=' * 70}")
        run_per_ion_section(all_ions, total_dropped, ion_out, args, fs)
        print(f"\n  Per-ion outputs in {ion_out}/")

    print("\n✓ Done.")


if __name__ == "__main__":
    main()

'''
# one channel: per-run outputs + per-ion boxplots
python3 aggregate_permeation_ion_stats.py \
  /media/ziyue/328bfc27-c1a6-4fce-9199-95c389ecd48d/Konstantina/girk_analyser_results/G12

# both channels in one call: per-ion boxplots compare them, with tests
python3 aggregate_permeation_ion_stats.py \
  /media/ziyue/328bfc27-c1a6-4fce-9199-95c389ecd48d/Konstantina/girk_analyser_results/G12 \
  /media/ziyue/328bfc27-c1a6-4fce-9199-95c389ecd48d/Konstantina/girk_analyser_results/G2 \
  -stats

# per-ion residence boxes one per run
python3 aggregate_permeation_ion_stats.py <base_dir> -by run

# old behaviour only
python3 aggregate_permeation_ion_stats.py <base_dir> -noion
'''
'''
python3 aggregate_permeation_ion_stats.py \
  /media/ziyue/328bfc27-c1a6-4fce-9199-95c389ecd48d/Konstantina/girk_analyser_results/G12

python3 aggregate_permeation_ion_stats.py \
  /media/ziyue/328bfc27-c1a6-4fce-9199-95c389ecd48d/Konstantina/girk_analyser_results/G2

python3 aggregate_permeation_ion_stats.py \
  /media/ziyue/328bfc27-c1a6-4fce-9199-95c389ecd48d/Konstantina/girk_analyser_results/G12_D173N

python3 aggregate_permeation_ion_stats.py \
  /media/ziyue/328bfc27-c1a6-4fce-9199-95c389ecd48d/Konstantina/girk_analyser_results/G12_E141Q
'''